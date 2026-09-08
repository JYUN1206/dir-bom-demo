"""输出采购清单 Excel 与 reject 清单 Excel(rules.md / 手册 3.6)。

个人项目 / 模拟数据 —— 简化业务模型，真实拆单规则以贵司工艺为准。
purchase_list.xlsx 含「板材 / 五金 / 汇总」三页;reject_list.xlsx 含原因列。
成本口径与 prices.xlsx 一致:板材按张、五金按个/副/套/包。
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.bom_engine import BoardPart
from src.hardware_calc import HardwareItem
from src.optimizer import MaterialLine
from src.parser import Order
from src.validator import RejectRecord

HEADER_FONT = Font(bold=True)


@dataclass(frozen=True, slots=True)
class OrderResult:
  """一个合规订单的完整计算结果(供报表使用)。"""

  order: Order
  parts: list[BoardPart]
  hardware: list[HardwareItem]
  plan: list[MaterialLine]


def _sheet_style(sheet, headers, rows):
  sheet.append(headers)
  for cell in sheet[1]:
    cell.font = HEADER_FONT
  for row in rows:
    sheet.append(row)
  # 列宽自适应 = max(表头,该列最长值) + 2
  max_widths = [len(str(h)) for h in headers]
  for row in rows:
    for col_idx, value in enumerate(row):
      max_widths[col_idx] = max(max_widths[col_idx], len(str(value)))
  for col_idx, width in enumerate(max_widths, start=1):
    sheet.column_dimensions[get_column_letter(col_idx)].width = width + 2
  sheet.freeze_panes = 'A2'


def hardware_cost(items, prices):
  return sum(item.quantity * prices[item.name]['price'] for item in items)


def totals(result, prices):
  """返回 (板材成本, 五金成本, 总成本)。"""
  board_cost = sum(line.cost for line in result.plan)
  hw_cost = round(hardware_cost(result.hardware, prices), 2)
  return round(board_cost, 2), hw_cost, round(board_cost + hw_cost, 2)


def build_purchase_workbook(results, prices):
  """构造采购清单 Workbook(板材/五金/汇总三页)。"""
  book = Workbook()

  board_sheet = book.active
  board_sheet.title = '板材'
  board_rows, hw_rows, summary_rows = [], [], []

  for result in results:
    order = result.order
    for line in result.plan:
      board_rows.append(
        [order.order_id, order.board_material if line.category == '主材' else '背板9mm',
         line.category, line.sheets, line.unit_price, line.cost]
      )
    for item in result.hardware:
      unit_price = prices[item.name]['price']
      hw_rows.append(
        [order.order_id, item.name, item.quantity, item.unit, unit_price,
         round(item.quantity * unit_price, 2)]
      )
    board_cost, hw_cost, total = totals(result, prices)
    summary_rows.append(
      [order.order_id, order.customer, order.cabinet_type, order.board_material,
       board_cost, hw_cost, total]
    )

  _sheet_style(board_sheet,
               ['订单号', '板材料号', '类型', '张数', '单价(元)', '成本(元)'], board_rows)
  hw_sheet = book.create_sheet('五金')
  _sheet_style(hw_sheet,
               ['订单号', '品名', '数量', '单位', '单价(元)', '成本(元)'], hw_rows)
  summary_sheet = book.create_sheet('汇总')
  _sheet_style(summary_sheet,
               ['订单号', '客户', '柜体类型', '主材', '板材成本(元)', '五金成本(元)', '总成本(元)'],
               summary_rows)
  return book


def build_reject_workbook(rejects):
  """构造 reject 清单 Workbook。"""
  book = Workbook()
  sheet = book.active
  sheet.title = 'reject'
  _sheet_style(sheet,
               ['订单号', '客户', '柜体类型', '违规部件', '方向', '实际值(mm)',
                '上限值(mm)', '原因'],
               [[r.order_id, r.customer, r.cabinet_type, r.part_name, r.direction,
                 r.value_mm, r.limit_mm, r.reason] for r in rejects])
  return book


def write_excel(path, workbook):
  path = Path(path)
  path.parent.mkdir(parents=True, exist_ok=True)
  workbook.save(path)
  return path


def write_reports(outdir, results, rejects, prices):
  """写 purchase_list.xlsx + reject_list.xlsx 到 outdir, 返回两个路径。"""
  outdir = Path(outdir)
  purchase_path = write_excel(outdir / 'purchase_list.xlsx',
                              build_purchase_workbook(results, prices))
  reject_path = write_excel(outdir / 'reject_list.xlsx',
                            build_reject_workbook(rejects))
  return purchase_path, reject_path


def readable(path):
  """校验生成的 Excel 可被重新打开。"""
  return load_workbook(Path(path)) is not None
